"""Excel generator for the Presupuestador Sage 50."""

import datetime
import io

from typing import cast

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from models import agrupar_lineas, subtotal_grupo


def generar_excel(datos_cliente, lineas, totales, notas="", agrupacion=None):
    """Generate an editable quote workbook and return its bytes."""
    wb = Workbook()
    ws = cast(Worksheet, wb.active)
    ws.title = "Presupuesto"
    hoy = datetime.date.today().strftime("%d/%m/%Y")
    num = datetime.datetime.now().strftime("PRE-%Y%m%d-%H%M")
    titulo = Font(name="Calibri", size=16, bold=True, color="1A5276")
    subtitulo = Font(name="Calibri", size=11, bold=True, color="2C3E50")
    normal = Font(name="Calibri", size=10)
    bold = Font(name="Calibri", size=10, bold=True)
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1A5276", end_color="1A5276", fill_type="solid")
    total_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    total_fill = PatternFill(start_color="1A5276", end_color="1A5276", fill_type="solid")
    group_fill = PatternFill(start_color="D9EAF7", end_color="D9EAF7", fill_type="solid")
    alt_fill = PatternFill(start_color="F2F7FB", end_color="F2F7FB", fill_type="solid")
    side = Side(style="thin", color="DEE2E6")
    border = Border(left=side, right=side, top=side, bottom=side)
    money = '#,##0.00 €'
    ws.merge_cells("A1:H1")
    ws["A1"] = "PRESUPUESTO"
    ws["A1"].font = titulo
    ws.merge_cells("A2:H2")
    ws["A2"] = f"Nº {num}  |  Fecha: {hoy}  |  Validez: {datos_cliente.get('validez', '30 días')}"
    ws["A2"].font = Font(name="Calibri", size=9, color="666666")
    row = 4
    ws.merge_cells(f"A{row}:H{row}")
    ws[f"A{row}"] = "DATOS DEL CLIENTE"
    ws[f"A{row}"].font = subtitulo
    row = 5
    for label, key in [("Empresa", "empresa"), ("CIF/NIF", "cif"), ("Contacto", "contacto"), ("Email", "email")]:
        if datos_cliente.get(key, ""):
            ws[f"A{row}"] = f"{label}:"
            ws[f"A{row}"].font = bold
            ws.merge_cells(f"B{row}:D{row}")
            ws[f"B{row}"] = datos_cliente[key]
            row += 1
    if datos_cliente.get("condiciones", ""):
        ws[f"A{row}"] = "Condiciones:"
        ws[f"A{row}"].font = bold
        ws.merge_cells(f"B{row}:D{row}")
        ws[f"B{row}"] = datos_cliente["condiciones"]
        row += 1
    row += 1
    ws.merge_cells(f"A{row}:H{row}")
    ws[f"A{row}"] = "DETALLE DEL PRESUPUESTO"
    ws[f"A{row}"].font = subtitulo
    row += 1
    headers = ["Código", "Descripción", "Plan", "Periodo", "Uds", "Precio Ud.", "Dto %", "Total"]
    widths = [16, 38, 14, 14, 8, 14, 10, 14]
    for col, (header, width) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = border
        ws.column_dimensions[get_column_letter(col)].width = width
    header_row = row
    row += 1
    for grupo, group_lines in agrupar_lineas(lineas, agrupacion):
        if grupo is not None:
            ws.merge_cells(f"A{row}:H{row}")
            cell = ws[f"A{row}"]
            cell.value = f"{'PLAN' if agrupacion == 'plan' else 'PERIODO'}: {grupo}"
            cell.font = bold
            cell.fill = group_fill
            row += 1
        for index, linea in enumerate(group_lines):
            values = [linea.get("codigo", ""), linea["descripcion"], linea.get("plan", "-"), linea.get("periodicidad", "-"), linea["cantidad"], linea["precio_unitario"], linea.get("dto_total_pct", 0), linea["total"]]
            for col, value in enumerate(values, 1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.font = normal
                cell.border = border
                if col >= 5:
                    cell.alignment = Alignment(horizontal="right")
                if col in (6, 8):
                    cell.number_format = money
                if col == 7:
                    cell.number_format = '0%'
                if index % 2 == 1:
                    cell.fill = alt_fill
            row += 1
        if grupo is not None:
            ws.cell(row=row, column=7, value=f"Subtotal {grupo}").font = bold
            cell = ws.cell(row=row, column=8, value=subtotal_grupo(group_lines))
            cell.font = bold
            cell.number_format = money
            row += 1
    data_end = row - 1
    for label, value, is_total in [("SUBTOTAL", totales["subtotal"], False), ("TOTAL + IVA", totales["total"], True)]:
        label_cell = ws.cell(row=row, column=7, value=label)
        value_cell = ws.cell(row=row, column=8, value=value)
        value_cell.number_format = money
        label_cell.alignment = value_cell.alignment = Alignment(horizontal="right")
        label_cell.font = value_cell.font = total_font if is_total else bold
        if is_total:
            label_cell.fill = value_cell.fill = total_fill
        row += 1
    if notas:
        row += 1
        ws.merge_cells(f"A{row}:H{row}")
        ws[f"A{row}"] = "OBSERVACIONES"
        ws[f"A{row}"].font = subtitulo
        row += 1
        ws.merge_cells(f"A{row}:H{row}")
        ws[f"A{row}"] = notas
    row += 2
    ws.merge_cells(f"A{row}:H{row}")
    ws[f"A{row}"] = f"Generado el {hoy} — alfonsoautomatiza.com — Cádiz, España"
    ws[f"A{row}"].font = Font(name="Calibri", size=8, color="999999")
    ws.auto_filter.ref = f"A{header_row}:H{data_end}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
